from pathlib import Path
from openpyxl import load_workbook

files = [
    Path("Network Device Inventory - Final Oct 21-2026.xlsx"),
    Path("Server-MasterSheet 2025 - reorganized.xlsx"),
    Path("Telcom Inventory_21st Oct 2026_reordered.xlsx"),
]

for path in files:
    print(f"\n### {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        print(f"\n[{sheet.title}]")
        shown = 0
        for row in sheet.iter_rows(values_only=True):
            if any(value is not None for value in row):
                print(row)
                shown += 1
                if shown == 8:
                    break
    workbook.close()